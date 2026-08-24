# -*- coding: utf-8 -*-
# flake8: noqa: E501
# pylint: disable=line-too-long
"""Document edit / create tools (self-built, Apache-2.0 clean).

[hanbao modification] These tools replace the Anthropic-proprietary
`docx` / `xlsx` skills that were removed for licensing reasons (I-019).
Instead of shipping third-party skill code, we build a minimal,
self-contained implementation on top of `python-docx` (MIT) and
`openpyxl` (MIT) — both permissive licenses compatible with
redistribution, unlike the removed Anthropic skills.

Scope is intentionally "简化版":
- create_docx / create_xlsx: build a document from plain text input.
- edit_docx / edit_xlsx: append content or replace/find text / write cells.
No styling engines, no template system — just enough for an agent to
produce and tweak Office documents on behalf of the user.
"""

import logging
import os
from pathlib import Path

from agentscope.message import TextBlock, ToolResultState
from agentscope.tool import ToolChunk

from .file_io import _resolve_file_path
from ...runtime.tool_registry import tool_descriptor
from ...utils.io_utils import get_path_lock

logger = logging.getLogger(__name__)


def _ensure_parent_dir(file_path: str) -> None:
    """Create the parent directory of *file_path* if it does not exist."""
    parent = os.path.dirname(os.path.abspath(file_path))
    if parent:
        os.makedirs(parent, exist_ok=True)


def _parse_rows(rows_text: str) -> list[list[str]]:
    """Parse tab-separated rows text into a list of cell lists.

    Each non-empty line becomes one row; cells are split on tabs.
    """
    rows: list[list[str]] = []
    for line in (rows_text or "").split("\n"):
        if line.strip() == "":
            continue
        rows.append([c for c in line.split("\t")])
    return rows


def _ok(text: str) -> ToolChunk:
    return ToolChunk(
        is_last=True,
        state=ToolResultState.SUCCESS,
        content=[TextBlock(type="text", text=text)],
    )


def _err(text: str) -> ToolChunk:
    return ToolChunk(
        is_last=True,
        state=ToolResultState.ERROR,
        content=[TextBlock(type="text", text=text)],
    )


@tool_descriptor(
    requires_sandbox=("file_write",),
    async_execution=True,
    tool_type="file",
    target_param="file_path",
    policy_name="CreateDocx",
    ui_description="Create a Word (.docx) document from text",
    ui_icon="📝",
)
async def create_docx(
    file_path: str,
    content: str,
    title: str = "",
) -> ToolChunk:
    """Create a new Word document (.docx) from plain text.

    Each non-empty line of *content* becomes its own paragraph. An optional
    *title* is added as the document title heading.

    Args:
        file_path (`str`): Output path. Relative paths resolve from the
            current workspace.
        content (`str`): Paragraph text, one paragraph per non-empty line.
        title (`str`, optional): Document title heading.
    """
    if not file_path:
        return _err("Error: No `file_path` provided.")
    try:
        from docx import Document  # lazy: MIT dep, graceful if missing
    except ImportError as exc:  # pragma: no cover - env guard
        return _err(f"Error: python-docx is not installed ({exc}).")

    resolved = _resolve_file_path(file_path)
    try:
        _ensure_parent_dir(resolved)
        doc = Document()
        if title:
            doc.add_heading(title, level=0)
        for line in (content or "").split("\n"):
            if line.strip() == "":
                continue
            doc.add_paragraph(line)
        doc.save(resolved)
        return _ok(f"Created Word document at {resolved}.")
    except Exception as exc:
        logger.warning(f"create_docx failed: {exc}")
        return _err(f"Error: create_docx failed: {exc}")


@tool_descriptor(
    requires_sandbox=("file_write",),
    async_execution=True,
    tool_type="file",
    target_param="file_path",
    policy_name="EditDocx",
    ui_description="Edit a Word (.docx) document: append or replace text",
    ui_icon="📝",
)
async def edit_docx(
    file_path: str,
    text: str = "",
    old_text: str = "",
    new_text: str = "",
    operation: str = "append",
) -> ToolChunk:
    """Edit an existing Word document (.docx).

    - operation="append": append *text* as new paragraphs at the end.
    - operation="replace": replace every occurrence of *old_text* with
      *new_text* within a paragraph (multi-paragraph spans are not handled
      in this simplified version).

    Args:
        file_path (`str`): Path to the existing .docx.
        text (`str`, optional): Text to append (append mode).
        old_text (`str`, optional): Text to find (replace mode).
        new_text (`str`, optional): Replacement text (replace mode).
        operation (`str`): "append" (default) or "replace".
    """
    if not file_path:
        return _err("Error: No `file_path` provided.")
    try:
        from docx import Document  # lazy: MIT dep, graceful if missing
    except ImportError as exc:  # pragma: no cover - env guard
        return _err(f"Error: python-docx is not installed ({exc}).")

    resolved = _resolve_file_path(file_path)
    if not os.path.isfile(resolved):
        return _err(f"Error: The file {resolved} does not exist.")

    try:
        async with get_path_lock(resolved):
            doc = Document(resolved)
            if operation == "replace":
                if not old_text:
                    return _err("Error: replace mode requires `old_text`.")
                hits = 0
                for para in doc.paragraphs:
                    if old_text in para.text:
                        para.text = para.text.replace(old_text, new_text or "")
                        hits += 1
                if hits == 0:
                    return _err(
                        f"Error: `{old_text}` not found in {file_path}.",
                    )
                doc.save(resolved)
                return _ok(
                    f"Replaced `{old_text}` in {hits} paragraph(s) of "
                    f"{file_path}.",
                )
            # default: append
            for line in (text or "").split("\n"):
                if line.strip() == "":
                    continue
                doc.add_paragraph(line)
            doc.save(resolved)
            return _ok(f"Appended content to {file_path}.")
    except Exception as exc:
        logger.warning(f"edit_docx failed: {exc}")
        return _err(f"Error: edit_docx failed: {exc}")


@tool_descriptor(
    requires_sandbox=("file_write",),
    async_execution=True,
    tool_type="file",
    target_param="file_path",
    policy_name="CreateXlsx",
    ui_description="Create an Excel (.xlsx) spreadsheet from tab-separated rows",
    ui_icon="📊",
)
async def create_xlsx(
    file_path: str,
    rows: str,
    sheet_name: str = "Sheet1",
) -> ToolChunk:
    """Create a new Excel spreadsheet (.xlsx).

    *rows* is text where each non-empty line is a row and cells within a
    row are separated by TAB characters.

    Args:
        file_path (`str`): Output path. Relative paths resolve from the
            current workspace.
        rows (`str`): Tab-separated rows, one row per non-empty line.
        sheet_name (`str`, optional): Name of the first worksheet.
    """
    if not file_path:
        return _err("Error: No `file_path` provided.")
    try:
        from openpyxl import Workbook  # lazy: MIT dep, graceful if missing
    except ImportError as exc:  # pragma: no cover - env guard
        return _err(f"Error: openpyxl is not installed ({exc}).")

    resolved = _resolve_file_path(file_path)
    try:
        _ensure_parent_dir(resolved)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name or "Sheet1"
        for row in _parse_rows(rows):
            ws.append(row)
        wb.save(resolved)
        return _ok(f"Created Excel spreadsheet at {resolved}.")
    except Exception as exc:
        logger.warning(f"create_xlsx failed: {exc}")
        return _err(f"Error: create_xlsx failed: {exc}")


@tool_descriptor(
    requires_sandbox=("file_write",),
    async_execution=True,
    tool_type="file",
    target_param="file_path",
    policy_name="EditXlsx",
    ui_description="Edit an Excel (.xlsx): append a row or write a cell",
    ui_icon="📊",
)
async def edit_xlsx(
    file_path: str,
    operation: str = "append_row",
    row: str = "",
    cell: str = "",
    value: str = "",
    sheet_name: str = "",
) -> ToolChunk:
    """Edit an existing Excel spreadsheet (.xlsx).

    - operation="append_row": append *row* (tab-separated cells) to the
      sheet.
    - operation="write_cell": write *value* into *cell* (e.g. "A1" or
      "Sheet2!B3").

    Args:
        file_path (`str`): Path to the existing .xlsx.
        operation (`str`): "append_row" (default) or "write_cell".
        row (`str`, optional): Tab-separated cells to append.
        cell (`str`, optional): Cell reference for write_cell.
        value (`str`, optional): Value to write for write_cell.
        sheet_name (`str`, optional): Target sheet (defaults to active).
    """
    if not file_path:
        return _err("Error: No `file_path` provided.")
    try:
        from openpyxl import load_workbook  # lazy: MIT dep
    except ImportError as exc:  # pragma: no cover - env guard
        return _err(f"Error: openpyxl is not installed ({exc}).")

    resolved = _resolve_file_path(file_path)
    if not os.path.isfile(resolved):
        return _err(f"Error: The file {resolved} does not exist.")

    try:
        async with get_path_lock(resolved):
            wb = load_workbook(resolved)
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.active

            if operation == "write_cell":
                if not cell:
                    return _err("Error: write_cell requires `cell`.")
                target = cell
                if "!" in target:
                    sheet_part, target = target.split("!", 1)
                    if sheet_part in wb.sheetnames:
                        ws = wb[sheet_part]
                ws[target] = value
                wb.save(resolved)
                return _ok(f"Wrote value to {cell} in {file_path}.")
            # default: append_row
            cells = [c for c in (row or "").split("\t")]
            ws.append(cells)
            wb.save(resolved)
            return _ok(f"Appended row to {file_path}.")
    except Exception as exc:
        logger.warning(f"edit_xlsx failed: {exc}")
        return _err(f"Error: edit_xlsx failed: {exc}")
